# -*- coding: utf-8 -*-
"""Clona una hoja del EJEMPLO de la catedra al presupuesto de ShopMetrics.

    python3 scripts/clonar_hoja_ejemplo.py "Estructura"

Para las hojas cuyo contenido es de la catedra y no del proyecto --las escalas
de probabilidad e impacto, la interpretacion del nivel de riesgo-- no hay nada
que adaptar: se copian tal cual, con su formato, que es justamente lo que el
profesor corrige.

Se copia atributo por atributo y no `_style`, porque ese atributo guarda
indices a las tablas del libro de origen y en otro libro apuntan a cualquier
otra cosa.
"""
from __future__ import annotations

import os
import sys
from copy import copy

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(os.path.dirname(AQUI), "documento")
LIBRO = os.path.join(DOC, "Presupuesto financiero ShopMetrics.xlsx")
EJEMPLO = os.path.join(DOC, "Presupuesto financiero EJEMPLO V1.xlsx")


def clonar(origen, destino, con_valores=True, filas=None, columnas=None):
    filas = filas or origen.max_row
    columnas = columnas or origen.max_column

    for rango in list(destino.merged_cells.ranges):
        destino.unmerge_cells(str(rango))
    for fila in destino.iter_rows(min_row=1, max_row=destino.max_row,
                                  max_col=destino.max_column):
        for celda in fila:
            celda.value = None

    for fila in origen.iter_rows(min_row=1, max_row=filas, max_col=columnas):
        for celda in fila:
            nueva = destino.cell(celda.row, celda.column)
            if con_valores:
                nueva.value = celda.value
            if not celda.has_style:
                continue
            nueva.font = copy(celda.font)
            nueva.fill = copy(celda.fill)
            nueva.border = copy(celda.border)
            nueva.alignment = copy(celda.alignment)
            nueva.number_format = celda.number_format
            nueva.protection = copy(celda.protection)

    for letra, dim in origen.column_dimensions.items():
        d = destino.column_dimensions[letra]
        d.width, d.hidden, d.bestFit = dim.width, dim.hidden, dim.bestFit
    for idx, dim in origen.row_dimensions.items():
        if idx <= filas and dim.height:
            destino.row_dimensions[idx].height = dim.height
    for rango in list(origen.merged_cells.ranges):
        if rango.max_row <= filas and rango.max_col <= columnas:
            destino.merge_cells(str(rango))

    destino.sheet_view.showGridLines = origen.sheet_view.showGridLines
    if origen.freeze_panes:
        destino.freeze_panes = origen.freeze_panes


def main() -> int:
    nombres = sys.argv[1:] or ["Estructura"]
    ejemplo = openpyxl.load_workbook(EJEMPLO)
    libro = openpyxl.load_workbook(LIBRO)
    for n in nombres:
        if n not in ejemplo.sheetnames or n not in libro.sheetnames:
            print("  %s: no está en los dos libros" % n); continue
        clonar(ejemplo[n], libro[n])
        llenas = sum(1 for f in libro[n].iter_rows() for c in f if c.value not in (None, ""))
        print("  %-34s %d celdas" % (n, llenas))
    libro.save(LIBRO)
    print("guardado")
    return 0


if __name__ == "__main__":
    sys.exit(main())
