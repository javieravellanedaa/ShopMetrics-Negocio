#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Rehace la hoja 'Costos RRHH' clonando el formato de Lucas Fraguaga.

No reconstruye el formato "a mano" ni lo toma del template: clona la hoja
'Costos RRHH' de 'PresupuestoFinanciero - Lucas Fraguaga.xlsx' --fuentes,
bordes, rellenos, formatos numericos, anchos, altos y combinaciones-- y
despues escribe nuestros datos encima, conservando su estructura:

  * cuadro 'Costo mensual por puesto' con la columna Area COMBINADA en
    vertical por area, sin filas vacias entre grupos;
  * encabezado de dos filas, con la alicuota de cada concepto debajo de su
    nombre (no una fila suelta de 'Alicuotas patronales');
  * las siete columnas de aportes de Lucas, incluida AAFF;
  * un unico bloque por ejercicio, con Cantidad y Costo intercalados por mes,
    'Junio + SAC' y 'Diciembre + SAC' y Total anual;
  * sin leyenda de colores ni puestos pintados: eso era del template.

Lo unico que agregamos sobre el formato de Lucas es una columna de
observaciones a la derecha del cuadro, porque Scali pidio expresamente saber
quien cubre cada puesto y en que hoja se imputa su costo (35:12).
"""
import os
import sys
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter as L

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARCHIVO = os.path.join(RAIZ, "documento", "Presupuesto financiero ShopMetrics.xlsx")
LUCAS = os.path.join(RAIZ, "PresupuestoFinanciero - Lucas Fraguaga.xlsx")

DEP, CUB, TER = "dep", "cub", "ter"

# (area, puesto, estado, sueldo bruto, observacion)
PUESTOS = [
    ("Gerencia", "Gerente General (Fundador / CEO)", DEP, 2000,
     "En relación de dependencia"),
    ("Gerencia", "Gerente de Sistemas (Fundador / CTO)", DEP, 2000,
     "En relación de dependencia"),
    ("Administración", "Administrativo contable y de pagos", CUB, 900,
     "Lo cubre el Gerente General"),
    ("Administración", "Liquidación de sueldos e impuestos", TER, 800,
     "Estudio contable, abono mensual en costos fijos"),
    ("Administración", "Administrativo de cobranzas", CUB, 850,
     "Lo cubre el Gerente General"),
    ("Comercial", "Ejecutivo comercial de altas", TER, 1000,
     "Freelance por comisión sobre cada alta, en costos variables"),
    ("Comercial", "Atención al cliente y posventa", CUB, 900,
     "Lo cubre el Gerente General"),
    ("Marketing", "Gestión de campañas y contenidos", TER, 1100,
     "Agencia de marketing, abono mensual en costos fijos"),
    ("Sistemas", "Arquitecto de software", TER, 2200,
     "Contratado por proyecto; el desarrollo inicial es inversión"),
    ("Sistemas", "Desarrollador", TER, 1800,
     "Mantenimiento evolutivo contratado, en costos fijos"),
    ("Sistemas", "Científico de datos", CUB, 2000,
     "Lo cubre el Gerente de Sistemas"),
    ("Sistemas", "DevOps / infraestructura cloud", CUB, 1900,
     "Lo cubre el Gerente de Sistemas"),
    ("Servicio técnico", "Técnico de instalación y soporte", TER, 1100,
     "Freelance por hora según el Anexo, en costos variables"),
    ("Servicio técnico", "Operador de monitoreo de sensores", CUB, 900,
     "Lo cubre el Gerente de Sistemas"),
]

UNO = [1] * 12
DOTACION = {0: {a: UNO for a in (2026, 2027, 2028)},      # Gerente General
            1: {a: UNO for a in (2026, 2027, 2028)}}      # Gerente de Sistemas

# los siete conceptos de Lucas; el seguro de vida es un monto fijo por persona
APORTES = [("Jubilación", 0.1235, "tasa"), ("Ley 19032", 0.0157, "tasa"),
           ("Obra social", 0.06, "tasa"), ("FNE", 0.0108, "tasa"),
           ("AAFF", 0.054, "tasa"), ("S. de Vida", 1.20, "fijo"),
           ("ART (Aprox.)", 0.03, "tasa")]

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio + SAC", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre + SAC"]
ANIOS = (2026, 2027, 2028)

CANT = [3 + 2 * i for i in range(12)]        # C,E,G,...,Y
COSTO = [4 + 2 * i for i in range(12)]       # D,F,H,...,Z
TOT = 27                                     # AA
OBS = 14                                     # N: observaciones nuestras


def clonar(src, dst):
    for k, d in src.column_dimensions.items():
        dst.column_dimensions[k].width = d.width
        dst.column_dimensions[k].hidden = d.hidden
    for k, d in src.row_dimensions.items():
        dst.row_dimensions[k].height = d.height
    for row in src.iter_rows():
        for c in row:
            n = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                n.font = copy(c.font); n.fill = copy(c.fill); n.border = copy(c.border)
                n.alignment = copy(c.alignment); n.number_format = c.number_format
                n.protection = copy(c.protection)
    for r in src.merged_cells.ranges:
        dst.merge_cells(str(r))
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_properties.tabColor = src.sheet_properties.tabColor
    dst.freeze_panes = src.freeze_panes


def main() -> int:
    wb = openpyxl.load_workbook(ARCHIVO)
    wl = openpyxl.load_workbook(LUCAS)
    fuente = wl["Costos RRHH"]

    idx = wb.sheetnames.index("Costos RRHH")
    wb.remove(wb["Costos RRHH"])
    r = wb.create_sheet("Costos RRHH", idx)
    clonar(fuente, r)

    # ---- estilos modelo, tomados del clon de Lucas ------------------------
    M = {}
    for nombre, ref in (("bloque", "B11"), ("enc", "C12"), ("alic", "E13"),
                        ("area", "B14"), ("puesto", "C14"), ("bruto", "D14"),
                        ("aporte", "E14"), ("fijo", "J14"), ("sumap", "L14"),
                        ("mensual", "M14"), ("tit_anio", "B32"), ("mes", "C33"),
                        ("subenc", "C34"), ("pto_anio", "B35"), ("cant", "C35"),
                        ("costo", "D35"), ("tot_rot", "B44"), ("tot_num", "C44")):
        c = fuente[ref]
        M[nombre] = {"font": copy(c.font), "fill": copy(c.fill), "border": copy(c.border),
                     "alignment": copy(c.alignment), "number_format": c.number_format}

    def poner(ref, valor, modelo):
        c = r[ref] if isinstance(ref, str) else ref
        c.value = valor
        m = M[modelo]
        c.font = copy(m["font"]); c.fill = copy(m["fill"]); c.border = copy(m["border"])
        c.alignment = copy(m["alignment"]); c.number_format = m["number_format"]
        return c

    # ---- limpieza de la fila 9 hacia abajo (la cabecera 1-7 se conserva) ---
    from openpyxl.styles import Border, Font, Alignment, PatternFill
    for rango in list(r.merged_cells.ranges):
        if rango.max_row >= 9:
            r.unmerge_cells(str(rango))
    for fila in r.iter_rows(min_row=9, max_row=max(r.max_row, 95), max_col=max(r.max_column, 31)):
        for c in fila:
            c.value = None
            c.fill = PatternFill(); c.border = Border(); c.font = Font()
            c.alignment = Alignment(); c.number_format = "General"

    # ---- cabecera: nuestros ejercicios ------------------------------------
    for i, anio in enumerate(ANIOS):
        for col in (2 + i, 8 + i):
            c = r.cell(row=5, column=col)
            c.value = anio
            c.number_format = "0"          # el anio es un rotulo, no un numero con miles
        c = r.cell(row=6, column=2 + i)
        c.value = "=Hipótesis!C%d" % (24 + i)
        c.number_format = "0.0%"           # la participacion lleva un decimal
        r.cell(row=7, column=2 + i).value = "=Hipótesis!D%d" % (24 + i)

    # ---- cuadro de puestos -------------------------------------------------
    F_ENC, F_ALI, F0 = 12, 13, 14
    poner("B11", "Costo mensual por puesto", "bloque")
    r.merge_cells("B11:L11")

    for col, t in ((2, "Area"), (3, "Puesto"), (4, "Sueldo bruto")):
        poner(r.cell(row=F_ENC, column=col), t, "enc")
        poner(r.cell(row=F_ALI, column=col), None, "enc")
        r.merge_cells(start_row=F_ENC, start_column=col, end_row=F_ALI, end_column=col)
    for i, (nombre, valor, tipo) in enumerate(APORTES):
        poner(r.cell(row=F_ENC, column=5 + i), nombre, "enc")
        c = poner(r.cell(row=F_ALI, column=5 + i), valor, "alic")
        c.number_format = '"$ "#,##0.00" c/u"' if tipo == "fijo" else "0.00%"
    for col, t in ((12, "Total aportes"), (13, "Costo mensual para empresa")):
        poner(r.cell(row=F_ENC, column=col), t, "enc")
        poner(r.cell(row=F_ALI, column=col), None, "enc")
        r.merge_cells(start_row=F_ENC, start_column=col, end_row=F_ALI, end_column=col)
    poner(r.cell(row=F_ENC, column=OBS), "Quién lo cubre y dónde se imputa el costo", "enc")
    r.merge_cells(start_row=F_ENC, start_column=OBS, end_row=F_ALI, end_column=TOT)

    fila_de = {}
    for i, (area, puesto, estado, bruto, obs) in enumerate(PUESTOS):
        f = F0 + i
        fila_de[i] = f
        poner("B%d" % f, None, "area")
        poner("C%d" % f, puesto, "puesto")
        poner("D%d" % f, bruto, "bruto")
        for j, (_, _, tipo) in enumerate(APORTES):
            col = L(5 + j)
            if tipo == "fijo":
                c = poner("%s%d" % (col, f), "=$%s$%d" % (col, F_ALI), "fijo")
            else:
                c = poner("%s%d" % (col, f), "=$D%d*%s$%d" % (f, col, F_ALI), "aporte")
        poner("L%d" % f, "=SUM(E%d:K%d)" % (f, f), "sumap")
        poner("M%d" % f, "=D%d+L%d" % (f, f), "mensual")
        c = poner(r.cell(row=f, column=OBS), obs, "puesto")
        c.alignment = copy(M["puesto"]["alignment"])
        r.merge_cells(start_row=f, start_column=OBS, end_row=f, end_column=TOT)

    # la columna Area va combinada en vertical por grupo, como en la hoja de Lucas
    i = 0
    while i < len(PUESTOS):
        j = i
        while j + 1 < len(PUESTOS) and PUESTOS[j + 1][0] == PUESTOS[i][0]:
            j += 1
        r["B%d" % (F0 + i)] = PUESTOS[i][0]
        if j > i:
            r.merge_cells(start_row=F0 + i, start_column=2, end_row=F0 + j, end_column=2)
        i = j + 1

    # ---- un bloque por ejercicio ------------------------------------------
    totales = {}
    f = F0 + len(PUESTOS) + 1
    for anio in ANIOS:
        poner("B%d" % f, "Costos de RRHH - Año %d" % anio, "tit_anio")
        r.merge_cells(start_row=f, start_column=2, end_row=f, end_column=TOT)
        f_enc, f_sub, f_ini = f + 1, f + 2, f + 3

        poner("B%d" % f_enc, "Puesto", "mes")
        poner("B%d" % f_sub, None, "mes")
        r.merge_cells("B%d:B%d" % (f_enc, f_sub))
        for i, mes in enumerate(MESES):
            poner(r.cell(row=f_enc, column=CANT[i]), mes, "mes")
            poner(r.cell(row=f_enc, column=COSTO[i]), None, "mes")
            r.merge_cells(start_row=f_enc, start_column=CANT[i],
                          end_row=f_enc, end_column=COSTO[i])
            poner(r.cell(row=f_sub, column=CANT[i]), "Cantidad", "subenc")
            poner(r.cell(row=f_sub, column=COSTO[i]), "Costo", "subenc")
        poner(r.cell(row=f_enc, column=TOT), "Total anual", "mes")
        poner(r.cell(row=f_sub, column=TOT), None, "mes")
        r.merge_cells(start_row=f_enc, start_column=TOT, end_row=f_sub, end_column=TOT)

        for i, (_, puesto, _, _, _) in enumerate(PUESTOS):
            fp = f_ini + i
            ref = "$M$%d" % fila_de[i]
            poner("B%d" % fp, puesto, "pto_anio")
            dot = DOTACION.get(i, {}).get(anio, [0] * 12)
            for m in range(12):
                poner(r.cell(row=fp, column=CANT[m]), dot[m], "cant")
                cant = "%s%d" % (L(CANT[m]), fp)
                if m in (5, 11):      # el SAC: medio sueldo mas, en junio y en diciembre
                    formula = "=%s*%s+%s%d/2" % (cant, ref, L(COSTO[m - 1]), fp)
                else:
                    formula = "=%s*%s" % (cant, ref)
                poner(r.cell(row=fp, column=COSTO[m]), formula, "costo")
            poner(r.cell(row=fp, column=TOT),
                  "=" + "+".join("%s%d" % (L(c), fp) for c in COSTO), "costo")

        f_tot = f_ini + len(PUESTOS)
        totales[anio] = f_tot
        poner("B%d" % f_tot, "Totales", "tot_rot")
        for col in CANT + COSTO + [TOT]:
            c = poner(r.cell(row=f_tot, column=col),
                      "=SUM(%s%d:%s%d)" % (L(col), f_ini, L(col), f_tot - 1), "tot_num")
            if col in CANT:
                c.number_format = "0"
        f = f_tot + 4

    for i, anio in enumerate(ANIOS):
        r.cell(row=6, column=8 + i).value = "=$%s$%d" % (L(TOT), totales[anio])

    r.column_dimensions["C"].width = 38
    r.column_dimensions[L(OBS)].width = 13.5

    wb.save(ARCHIVO)
    print("Costos RRHH clonada del formato de Lucas.")
    print("  puestos: %d en %d áreas (columna Área combinada por grupo)" % (
        len(PUESTOS), len({p[0] for p in PUESTOS})))
    print("  aportes: %s" % ", ".join(a[0] for a in APORTES))
    print("  bloques: %s" % ", ".join("%d→Totales f%d" % (a, totales[a]) for a in ANIOS))
    return 0


if __name__ == "__main__":
    sys.exit(main())
