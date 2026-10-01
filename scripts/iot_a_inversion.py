#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reconoce los dispositivos IoT como inversion y los saca de costos variables.

Scali, 35:40-35:51, revisando nuestras hojas:
    "vos tenes alguna solucion de IoT QUE NO SON INSUMOS, que los insumos son
     variables, eso lo vas a tener que calcular A PARTIR DE LAS INSTALACIONES
     QUE HAGAS"
y acto seguido, 36:05: "a ver el modelo de inversion".

Lucas Fraguaga lo resuelve igual: el stock de microcontroladores y sensores
figura en su Mod. inversion, ano por ano, y NO aparece en su hoja de
Amortizaciones, que solo lleva servidor, notebooks, celulares, vehiculo y
muebles. Los dispositivos no se amortizan: son stock que rota, se entrega en
comodato y vuelve a la empresa cuando el comercio se da de baja.

Qué hace:
  1. agrega a cada ejercicio del Mod. inversion las dos lineas de dispositivos,
     con la cantidad tomada por formula de las altas de Proy. ventas, de modo
     que la inversion se recalcula sola si cambia la proyeccion;
  2. limpia las filas de "Insumos IoT por instalacion" de Costos variables en
     los tres bloques, sin dejar filas en cero;
  3. no toca Amortizaciones.
"""
import os
import re
import sys
from copy import copy

import docx
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARCHIVO = os.path.join(RAIZ, "documento", "Presupuesto financiero ShopMetrics.xlsx")
DOCX = os.path.join(RAIZ, "documento", "STF_Gomez_Javier_E1_v1.docx")

# ejercicio -> (primera fila libre del bloque, fila de altas Vidriera, fila de altas Cadena)
# en 2028 se arranca en la 70 porque la 69 ya lleva la renovacion de notebooks
BLOQUES = {2026: (32, 20, 21), 2027: (50, 83, 84), 2028: (70, 145, 146)}

PRECIO_VIDRIERA, PRECIO_CADENA = 34, 170
DISPOSITIVOS = [
    ("Dispositivos IoT en comodato - Plan Vidriera (sensor de vidriera, contador "
     "de puerta, gateway y kit de montaje)", PRECIO_VIDRIERA, 1,
     "Una unidad por instalación del Plan Vidriera, según las altas de Proy. ventas. "
     "Se entrega en comodato, vuelve a la empresa en la baja y se reinstala: es stock "
     "que rota, no un bien de uso, y por eso no se amortiza."),
    ("Dispositivos IoT en comodato - Plan Cadena (cinco locales por alta)",
     PRECIO_CADENA, 2,
     "Un kit de cinco locales por instalación del Plan Cadena, según las altas de "
     "Proy. ventas. Mismo criterio: comodato reutilizable, no se amortiza."),
]

# por bloque de Costos variables: (filas de insumos IoT que salen, fila que
# hay que subir para que no quede un hueco en el medio del cuadro)
VARIABLES_FUERA = [((25, 26), 27), ((48, 49), 50), ((71, 72), 73)]


def word():
    """El punto 6 seguia poniendo el equipamiento entre los costos variables,
    y en el mismo parrafo decia que opera como capital de trabajo que rota."""
    nuevo = (
        "Los costos variables crecen de manera proporcional a la actividad y son los que determinan el "
        "margen de contribución por cliente. Se distinguen dos grupos. El primero se devenga por cada alta "
        "concretada: las horas de instalación del técnico, que el anexo de capacidad operativa estima en "
        "media hora para el Plan Básico, tres horas para el Plan Vidriera y doce horas para el Plan Cadena "
        "a un costo de siete dólares la hora, la comisión variable del vendedor por cada comercio "
        "incorporado y la reposición del equipamiento que no se recupera. El segundo grupo se devenga mes a "
        "mes por cada comercio activo: las horas de soporte y mantenimiento, estimadas en un cuarto de hora "
        "mensual por comercio, el consumo de infraestructura cloud según el volumen de datos procesados y "
        "las comisiones de los medios de pago sobre lo efectivamente cobrado. El equipamiento que se "
        "instala en el local no integra estos costos: se entrega en comodato, se recupera y se reacondiciona "
        "cuando el comercio se da de baja, de modo que opera como capital de trabajo que rota y no como un "
        "costo perdido, y por eso se reconoce en el modelo de inversión y no se amortiza. Esto resulta "
        "determinante en un segmento de rotación alta y abono bajo.")
    d = docx.Document(DOCX)
    for p in d.paragraphs:
        if "El primero se devenga por cada alta concretada" in p.text:
            p.runs[0].text = nuevo
            for r in p.runs[1:]:
                r.text = ""
            d.save(DOCX)
            return True
    return False


def main() -> int:
    if "--word" in sys.argv:
        print("punto 6 del informe: %s" % ("equipamiento fuera de costos variables" if word()
                                           else "NO SE ENCONTRÓ"))
        return 0
    wb = openpyxl.load_workbook(ARCHIVO)
    mi, cv = wb["Mod. inversión"], wb["Costos variables"]

    def estilo(desde, hasta):
        hasta.font = copy(desde.font); hasta.fill = copy(desde.fill)
        hasta.border = copy(desde.border); hasta.alignment = copy(desde.alignment)
        hasta.number_format = desde.number_format

    # modelos de formato: la linea de notebooks del ano cero
    modelo = {c: mi["%s12" % c] for c in "BCDEF"}

    agregadas = []
    for anio, (f0, fila_vid, fila_cad) in BLOQUES.items():
        for i, (concepto, precio, cual, nota) in enumerate(DISPOSITIVOS):
            f = f0 + i
            ref = fila_vid if cual == 1 else fila_cad
            mi["B%d" % f] = concepto
            mi["C%d" % f] = "='Proy. ventas'!C%d" % ref
            mi["D%d" % f] = precio
            mi["E%d" % f] = "=C%d*D%d" % (f, f)
            mi["F%d" % f] = nota
            for c in "BCDEF":
                estilo(modelo[c], mi["%s%d" % (c, f)])
        agregadas.append((anio, f0))

    # las filas que decian "No se preven inversiones en el ano" quedan reemplazadas
    # arriba; en 2028 la renovacion de notebooks sigue en su fila original

    def reubicar(formula, origen, destino):
        """Corre las referencias a la propia fila; deja intactas las externas."""
        def cambiar(m):
            if m.group(1):                       # 'Otra hoja'!B12 -> no se toca
                return m.group(0)
            if int(m.group(3)) != origen:
                return m.group(0)
            return "%s%d" % (m.group(2), destino)
        return re.sub(r"('[^']+'!)?(\$?[A-Z]{1,2}\$?)(\d+)", cambiar, formula)

    quitadas = 0
    for par, sube in VARIABLES_FUERA:
        destino = min(par)
        for c in cv[sube]:                        # la linea de abajo sube al hueco
            d = cv.cell(row=destino, column=c.column)
            v = c.value
            d.value = reubicar(v, sube, destino) if isinstance(v, str) and v.startswith("=") else v
            if c.has_style:
                d.font = copy(c.font); d.fill = copy(c.fill); d.border = copy(c.border)
                d.alignment = copy(c.alignment); d.number_format = c.number_format
        for f in (max(par), sube):                # quedan vacias al pie del bloque
            for c in cv[f]:
                c.value = None
            quitadas += 1

    wb.save(ARCHIVO)
    print("Dispositivos IoT reconocidos como inversión.")
    for anio, f in agregadas:
        print("   %d → Mod. inversión f%d y f%d" % (anio, f, f + 1))
    print("   %d filas liberadas en Costos variables, sin huecos en el cuadro" % quitadas)
    print("   Amortizaciones: sin cambios (los dispositivos no se amortizan)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
