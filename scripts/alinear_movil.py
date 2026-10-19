# -*- coding: utf-8 -*-
"""Alinea el informe con lo que el sistema es: una web responsive, no una app nativa.

El 10.4.3 y el 10.4.5 prometen una aplicacion movil desarrollada con React
Native para Android e iOS, con notificaciones push. Eso nunca se construyo y
no se va a construir: la decision fue que el acceso desde el telefono es el
mismo panel en un navegador, adaptado al ancho de la pantalla. Las tres
pantallas de telefono --CU-028-001, CU-029-001 y CU-030-001-- son esas mismas
vistas, y existen.

Dejarlo como estaba es peor que una imprecision: el profesor abre el sistema,
busca la app movil y no esta. Este script reemplaza las descripciones por lo
que el sistema hace, con el mismo nivel de detalle tecnico que tenian.

Tambien pone al dia tres titulos que quedaron con el nombre viejo de
CU-029-001 --la fuente unica ya dice "desde el movil"-- y la caratula, que
decia agosto.

    python3 scripts/alinear_movil.py
"""
from __future__ import annotations

import datetime
import os
import re
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORD = os.path.join(RAIZ, "documento", "STF_Gomez_Javier_E2_v2.docx")

MESES = ("Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre")

# Los parrafos que describen una app nativa. La clave es el texto que hay que
# encontrar --un prefijo alcanza-- y el valor es lo que va en su lugar.
REEMPLAZOS = [
    (
        "Aplicación móvil. App para Android e iOS",
        "Acceso desde el teléfono. El mismo panel web, adaptado al ancho de la "
        "pantalla del teléfono, para el personal operativo del centro: la bandeja "
        "de alertas ordenada por severidad, la ficha de cada alerta con su "
        "checklist de verificación y el registro de la resolución en el piso. No "
        "es una aplicación que se instala: se entra con el navegador, con la "
        "misma dirección y las mismas credenciales que desde una computadora.",
    ),
    (
        "Cliente móvil (operaciones, mantenimiento y seguridad). Aplicación móvil "
        "multiplataforma desarrollada con React Native",
        "Cliente móvil (operaciones, mantenimiento y seguridad). No hay un cliente "
        "aparte: es el mismo panel web servido por Next.js, con un diseño "
        "adaptable que reacomoda las pantallas cuando el ancho del navegador es el "
        "de un teléfono. Funciona en cualquier navegador moderno de Android o de "
        "iOS sin instalar nada. El aviso de una alerta crítica queda registrado en "
        "la entidad `notificacion` del modelo de datos y se ve al entrar; la "
        "notificación push al sistema operativo del teléfono requiere una "
        "aplicación instalada y queda fuera del alcance de esta entrega.",
    ),
    (
        "Frontend móvil. React Native como framework",
        "Frontend móvil. No hay un framework aparte para el teléfono: el panel es "
        "uno solo, hecho con Next.js y React, y se adapta con CSS al ancho de la "
        "pantalla. Mantener una sola base de código era el motivo por el que se "
        "había elegido React Native, y una web responsive lo consigue sin sumar "
        "una segunda aplicación que construir, firmar y publicar en dos tiendas.",
    ),
]

# Cambios de frase sueltos, dentro de parrafos que por lo demas estan bien.
FRASES = [
    ("el personal operativo (acceso a la aplicación móvil)",
     "el personal operativo (el mismo panel desde el navegador del teléfono)"),
    ("El operario recibe la notificación push en la aplicación móvil",
     "El operario ve el aviso de la alerta al entrar al panel desde su teléfono"),
    ("El operario abre una alerta asignada desde su aplicación móvil.",
     "El operario abre una alerta asignada desde el panel, en su teléfono."),
    # Las comillas de estos dos son rectas y no tipograficas: Word las deja
    # como se tipearon, y buscar con las curvas no los encuentra.
    ('El locatario abre la app móvil, accede a "Mi local"',
     'El locatario entra al portal desde su teléfono, accede a "Mi local"'),
    ('El locatario abre la app móvil y accede a "Mi local".',
     'El locatario entra al portal desde su teléfono y accede a "Mi local".'),
    # El nombre viejo de CU-029-001: la fuente unica ya dice "desde el movil".
    ("Gestionar alerta crítica desde la app móvil (Admin)",
     "Gestionar alerta crítica desde el móvil (Admin)"),
]

# En los capitulos de negocio la via de acceso desde el telefono se enumera
# como un canal digital mas, junto al dashboard y al portal. Ahi no hace falta
# describir la arquitectura, pero "aplicacion movil" al lado de "dashboard
# web" se lee como dos productos distintos, y es uno solo: el mismo panel en
# un navegador. "Acceso movil" dice lo mismo sin prometer una app que no hay.
FRASES_DE_NEGOCIO = [
    ("una aplicación móvil y un módulo", "un acceso móvil y un módulo"),
    ("dashboard web, aplicación móvil y portal de autogestión",
     "dashboard web, acceso móvil y portal de autogestión"),
    ("la aplicación móvil para el seguimiento diario",
     "el acceso móvil para el seguimiento diario"),
    ("la aplicación móvil canaliza la consulta cotidiana",
     "el acceso móvil canaliza la consulta cotidiana"),
]


def _escribir(parrafo: Paragraph, texto: str) -> None:
    """Deja el parrafo con un solo run, conservando el formato del primero."""
    if not parrafo.runs:
        parrafo.add_run(texto)
        return
    parrafo.runs[0].text = texto
    for r in parrafo.runs[1:]:
        r._element.getparent().remove(r._element)


def _reemplazar_frase(parrafo: Paragraph, viejo: str, nuevo: str) -> bool:
    """Cambia una frase dentro de un parrafo, conservando el resto.

    Se reescribe el parrafo entero en un solo run: la frase suele estar
    partida en varios runs --Word los corta donde quiere-- y buscarla run por
    run no la encuentra.
    """
    if viejo not in parrafo.text:
        return False
    _escribir(parrafo, parrafo.text.replace(viejo, nuevo))
    return True


def main() -> int:
    doc = Document(WORD)
    hechos = []

    for parrafo in doc.paragraphs:
        texto = parrafo.text.strip()
        for prefijo, nuevo in REEMPLAZOS:
            if texto.startswith(prefijo):
                _escribir(parrafo, nuevo)
                hechos.append("10.4  %s…" % prefijo[:46])
                break
        else:
            for viejo, nuevo in FRASES + FRASES_DE_NEGOCIO:
                if _reemplazar_frase(parrafo, viejo, nuevo):
                    hechos.append("frase  %s…" % viejo[:46])

    # El nombre viejo tambien esta en la fila del indice de 10.5.2.
    for tabla in doc.tables:
        if not tabla.rows or tabla.rows[0].cells[0].text.strip() != "ID":
            continue
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    for viejo, nuevo in FRASES:
                        if _reemplazar_frase(parrafo, viejo, nuevo):
                            hechos.append("10.5.2 %s…" % viejo[:40])

    # La caratula decia agosto.
    hoy = datetime.date.today()
    esperada = "Fecha: %d de %s %d" % (hoy.day, MESES[hoy.month - 1], hoy.year)
    for parrafo in doc.paragraphs[:60]:
        if re.match(r"^Fecha:\s*\d+ de \w+ \d{4}$", parrafo.text.strip()):
            if parrafo.text.strip() != esperada:
                _escribir(parrafo, esperada)
                hechos.append("caratula  %s" % esperada)
            break

    doc.save(WORD)
    for x in hechos:
        print("  " + x)
    print("cambios: %d" % len(hechos))
    return 0


if __name__ == "__main__":
    sys.exit(main())
