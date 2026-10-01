#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Nivela el punto 7 (estructura organizacional) con la hoja Costos RRHH.

Despues de rehacer Costos RRHH con el formato del template, la unica dotacion
en relacion de dependencia son los dos fundadores: la venta, la instalacion, el
soporte, el desarrollo, el marketing y la administracion contable se contratan.
Varios parrafos del punto 7 seguian describiendo Vendedores y Tecnicos propios
y una dotacion de diez personas al cierre.  Esto los pone al dia.
"""
import os
import sys

import docx

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(RAIZ, "documento", "STF_Gomez_Javier_E2_v2.docx")

# (fragmento que identifica el parrafo, texto nuevo completo)
CAMBIOS = [
    ("debilidad temporal a mitigar con incorporaciones planificadas en los meses 3-6",
     "Equipo fundador inicial reducido frente a la infraestructura gerencial de los competidores "
     "establecidos: debilidad temporal a mitigar con la contratación de especialistas a demanda "
     "—desarrollo, instalación, soporte, venta, marketing y administración contable— sin ampliar "
     "la estructura en relación de dependencia."),

    ("los Vendedores propios recorren el eje comercial",
     "La comercialización se apoya en dos canales complementarios. El canal principal es la venta "
     "personal directa, en la que ejecutivos comerciales contratados por comisión sobre cada alta "
     "concretada recorren el eje comercial y gestionan el ciclo completo con el comerciante, sin "
     "reventa ni distribuidores. A este se suma el canal institucional de las cámaras de "
     "comerciantes zonales y la CAME, que actúa como puerta de acceso y validación ante la "
     "comunidad de cada corredor. En el plano de la entrega del servicio, el acceso se materializa "
     "por tres vías digitales: el dashboard web, la aplicación móvil para el seguimiento diario "
     "desde el mostrador y el portal de autogestión, que además permite el alta remota del Plan "
     "Básico sin intervención de un técnico."),

    ("la dotación crece de forma escalonada y sólo cuando la proyección de ventas",
     "Lógica de crecimiento por etapas: la estructura se mantiene plana durante todo el horizonte "
     "del plan. En relación de dependencia están únicamente los dos socios fundadores, que ocupan "
     "la Gerencia General y la Gerencia de Sistemas y absorben entre ambos los puestos "
     "administrativos, comerciales y técnicos que no se tercerizan. El resto de las funciones se "
     "contrata a demanda y escala con el volumen sin exigir incorporaciones: la venta se paga por "
     "comisión sobre cada alta concretada, la instalación y el soporte se contratan por hora según "
     "lo que pide la cartera de cada mes, el mantenimiento evolutivo de la plataforma y las "
     "campañas se contratan por abono mensual, y la liquidación de sueldos e impuestos queda en un "
     "estudio contable. La hoja de costos de recursos humanos del presupuesto financiero despliega "
     "los catorce puestos de las seis áreas de la estructura y señala, para cada uno, quién lo "
     "cubre y en qué hoja se imputa su costo. El anexo de capacidad operativa verifica que el "
     "esquema alcanza: convierte las altas y los comercios activos de la proyección de ventas en "
     "horas hombre y muestra, mes a mes, cuántas horas de técnico freelance hay que contratar y "
     "qué holgura queda. La incorporación de personal en relación de dependencia se difiere al "
     "momento en que el volumen la justifique, por fuera del horizonte de tres años del plan."),

    ("de esa relación surge el momento en que corresponde incorporar un nuevo Técnico",
     "El área de Operaciones ejecuta la entrega física del servicio: la instalación de los "
     "dispositivos de conteo en los locales, el alta remota de los comercios del Plan Básico y el "
     "soporte posterior a los comercios activos. Es el área cuyo dimensionamiento está directamente "
     "atado a la proyección de ventas: el anexo de capacidad operativa del presupuesto financiero "
     "calcula las horas que consume cada alta según el plan (media hora en el Básico, tres horas en "
     "el Vidriera y doce horas en el Cadena) y las horas de soporte mensual por comercio activo "
     "(entre 0,15 y 0,60 horas según el plan), y de esa relación surge la cantidad de horas de "
     "técnico freelance que hay que contratar cada mes. El puesto no genera costo de recursos "
     "humanos: se imputa en costos variables, porque se paga por hora efectivamente trabajada y "
     "crece con la cartera."),

    ("La integran los Vendedores, que recorren los ejes comerciales asignados local por local",
     "El área Comercial es responsable de la captación de comercios y depende directamente del "
     "Gerente General, que absorbe su conducción durante todo el horizonte del plan. La integran "
     "ejecutivos comerciales contratados por comisión sobre cada alta concretada, que recorren los "
     "ejes comerciales asignados local por local. El esfuerzo comercial surge de la relación entre "
     "la cantidad de comercios que un ejecutivo puede visitar por mes, la tasa de conversión sobre "
     "esas visitas y el objetivo de altas de cada período; como la remuneración es por resultado, "
     "el costo se imputa en costos variables y no exige dotación en relación de dependencia."),
]


def reescribir(p, cuerpo):
    """Cambia el texto conservando el formato. Si el parrafo abre con un rotulo
    en negrita terminado en dos puntos, mantiene esa particion: el rotulo queda
    en el primer run y el cuerpo pasa a un segundo run sin negrita."""
    if not p.runs:
        return False
    rotulo = p.runs[0].bold and ":" in cuerpo[:60]
    if rotulo:
        corte = cuerpo.index(":") + 1
        p.runs[0].text = cuerpo[:corte]
        if len(p.runs) > 1:
            resto, extra = p.runs[1], p.runs[2:]
        else:
            resto, extra = p.add_run(), []
        resto.text = cuerpo[corte:]
        resto.bold = False
        resto.font.name = p.runs[0].font.name
        resto.font.size = p.runs[0].font.size
    else:
        p.runs[0].text = cuerpo
        extra = p.runs[1:]
    for r in extra:
        r.text = ""
    return True


def main() -> int:
    d = docx.Document(DOCX)
    hechos, faltan = [], []
    for marca, cuerpo in CAMBIOS:
        for p in d.paragraphs:
            if marca in p.text:
                if reescribir(p, cuerpo):
                    hechos.append(marca[:46])
                break
        else:
            faltan.append(marca[:46])
    d.save(DOCX)
    print("punto 7 nivelado: %d párrafos" % len(hechos))
    for h in hechos:
        print("   · %s…" % h)
    if faltan:
        print("NO ENCONTRADOS:")
        for f in faltan:
            print("   · %s…" % f)
    return 1 if faltan else 0


if __name__ == "__main__":
    sys.exit(main())
