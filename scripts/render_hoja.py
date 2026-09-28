# -*- coding: utf-8 -*-
"""Dibuja un rango de una hoja del presupuesto como imagen, calcando su formato.

Las figuras del capitulo 8 del informe son recortes de la planilla: mismos
rellenos, negritas y formatos de numero que en Excel. Cuando la planilla
cambia, hay que rehacerlas, y hacerlo a mano con capturas no es repetible.
Este modulo lee el rango con openpyxl --valores calculados y estilos-- y lo
dibuja con matplotlib imitando la grilla de Excel: anchos de columna, alturas
de fila, celdas combinadas, rellenos, bordes, alineacion y el formato
contable con el signo a la izquierda y el numero a la derecha.

    from render_hoja import rango
    rango("Mod. egresos", "A8:D16", "documento/img_punto8/06_egresos_tabla.png")
"""
from __future__ import annotations

import os
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARCHIVO = os.path.join(RAIZ, "documento", "Presupuesto financiero ShopMetrics.xlsx")
FUENTE = "Calibri"
for cand in ("Calibri", "Carlito", "Arial"):
    import matplotlib.font_manager as fm
    if cand in {f.name for f in fm.fontManager.ttflist}:
        FUENTE = cand; break

_wb_v = _wb_f = None


def libros():
    global _wb_v, _wb_f
    if _wb_v is None:
        _wb_v = openpyxl.load_workbook(ARCHIVO, data_only=True)
        _wb_f = openpyxl.load_workbook(ARCHIVO)
    return _wb_v, _wb_f


def es_ar(n, dec=2):
    s = "{:,.{d}f}".format(n, d=dec)
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def formatear(v, fmt):
    """Devuelve (texto, signo_izquierda) segun el formato de numero de Excel."""
    if v is None:
        return "", None
    if isinstance(v, str):
        return v, None
    if isinstance(v, bool):
        return str(v), None
    f = (fmt or "General")
    if "%" in f:
        dec = len(f.split(".")[1].replace("%", "")) if "." in f else 0
        return es_ar(v * 100, dec) + "%", None
    if "$" in f:
        dec = 2 if "0.00" in f else 0
        neg = v < 0
        txt = es_ar(abs(v), dec)
        return (("-" if neg else "") + txt), ("-$" if neg else "$")
    if "#,##0" in f or "0.00" in f or "0.0" in f:
        dec = len(f.split(".")[1].rstrip('"')) if "." in f and not f.endswith('"') else (2 if "0.00" in f else (1 if "0.0" in f else 0))
        dec = min(dec, 4)
        return es_ar(v, dec), None
    if isinstance(v, float) and abs(v - round(v)) > 1e-9:
        return es_ar(v, 2), None
    return es_ar(v, 0), None


def color(fill):
    try:
        if fill is None or fill.fill_type != "solid":
            return None
        c = fill.fgColor
        if c.type == "rgb" and c.rgb and c.rgb not in ("00000000",):
            rgb = c.rgb[-6:]
            return "#" + rgb
        if c.type == "theme":
            return {0: "#FFFFFF", 1: "#000000", 2: "#E7E6E6", 3: "#44546A", 4: "#4472C4",
                    5: "#ED7D31", 6: "#A5A5A5", 7: "#FFC000", 8: "#5B9BD5", 9: "#70AD47"}.get(c.theme, "#F2F2F2")
    except Exception:
        pass
    return None


def rango(hoja, celdas, salida, escala=1.0, min_col_px=40, quitar_vacias=False, columnas=None, filas_extra=None):
    v, f = libros()
    wsv, wsf = v[hoja], f[hoja]
    c0, r0, c1, r1 = range_boundaries(celdas)
    cols = list(columnas) if columnas else list(range(c0, c1 + 1))
    filas = list(range(r0, r1 + 1))
    if quitar_vacias:
        filas = [r for r in filas if any(wsv.cell(row=r, column=c).value not in (None, "") for c in cols)]
    if filas_extra:
        filas = sorted(set(filas) | set(filas_extra))

    # geometria en pixeles (aprox. de Excel a 96 dpi)
    # openpyxl guarda los anchos por rangos (min..max) bajo la primera letra:
    # hay que expandirlos, si no toda columna que no sea la primera del rango
    # aparece con el ancho por defecto
    ancho_col = {}
    for d in wsf.column_dimensions.values():
        if d.width:
            for cc in range(d.min or 1, (d.max or d.min or 1) + 1):
                ancho_col[cc] = d.width
    anchos = {}
    for c in cols:
        w = ancho_col.get(c, 8.43)
        anchos[c] = max(min_col_px, int(w * 7 + 5))
    def ancho_de(r, c):
        for m in wsf.merged_cells.ranges:
            if m.min_row == r and m.min_col == c:
                return sum(anchos.get(cc, int(ancho_col.get(cc, 8.43) * 7 + 5)) for cc in range(c, m.max_col + 1))
        return anchos.get(c, 60)

    altos = {}
    for r in filas:
        d = wsf.row_dimensions.get(r)
        h = d.height if d is not None and d.height else 15
        px = int(h * 96 / 72)
        if d is None or not d.height:              # Excel autoajusta: estimar por el texto envuelto
            for c in cols:
                cf = wsf.cell(row=r, column=c); cv = wsv.cell(row=r, column=c)
                if isinstance(cv.value, str) and (cf.alignment.wrap_text or len(cv.value) > 60):
                    tam = (cf.font.sz or 11)
                    chars = max(10, int(ancho_de(r, c) / (tam * 0.5)))
                    lineas = sum(max(1, -(-len(p) // chars)) for p in cv.value.split("\n"))
                    px = max(px, int(lineas * tam * 1.45) + 6)
        altos[r] = px

    # combinaciones dentro del rango
    combinadas = {}
    tapadas = set()
    for m in wsf.merged_cells.ranges:
        if m.min_row > r1 or m.max_row < r0 or m.min_col > c1 or m.max_col < c0:
            continue
        combinadas[(m.min_row, m.min_col)] = (m.max_row, m.max_col)
        for rr in range(m.min_row, m.max_row + 1):
            for cc in range(m.min_col, m.max_col + 1):
                if (rr, cc) != (m.min_row, m.min_col):
                    tapadas.add((rr, cc))

    W = sum(anchos.values()); H = sum(altos.values())
    fig = plt.figure(figsize=(W / 96 * escala, H / 96 * escala), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1]); ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")

    x_de = {}; x = 0
    for c in cols: x_de[c] = x; x += anchos[c]
    y_de = {}; y = 0
    for r in filas: y_de[r] = y; y += altos[r]

    for r in filas:
        for c in cols:
            if (r, c) in tapadas:
                continue
            cf = wsf.cell(row=r, column=c); cv = wsv.cell(row=r, column=c)
            r2, c2 = combinadas.get((r, c), (r, c))
            r2 = min(r2, r1); c2 = min(c2, c1)
            x0, y0 = x_de[c], y_de[r]
            # una combinacion puede abarcar filas o columnas que no se dibujan
            w = sum(anchos[cc] for cc in range(c, c2 + 1) if cc in anchos)
            h = sum(altos[rr] for rr in range(r, r2 + 1) if rr in altos)
            if w == 0 or h == 0:
                continue
            fondo = color(cf.fill)
            b = cf.border
            con_borde = any(getattr(b, lado).style for lado in ("left", "right", "top", "bottom")) if b else False
            ax.add_patch(Rectangle((x0, y0), w, h, facecolor=fondo or "white",
                                   edgecolor="#9A9A9A" if (con_borde or fondo) else "none", linewidth=0.6))
            texto, signo = formatear(cv.value, cf.number_format)
            if texto == "":
                continue
            if isinstance(texto, str):
                texto = texto.replace("$", r"\$")          # si no, matplotlib lo lee como formula
            fnt = cf.font
            tam = (fnt.sz or 11) * 1.0
            peso = "bold" if fnt.b else "normal"
            estilo = "italic" if fnt.i else "normal"
            col_txt = "#000000"
            try:
                if fnt.color is not None and fnt.color.type == "rgb" and fnt.color.rgb and fnt.color.rgb[-6:] not in ("000000",):
                    col_txt = "#" + fnt.color.rgb[-6:]
            except Exception:
                pass
            hor = (cf.alignment.horizontal or ("right" if isinstance(cv.value, (int, float)) and not isinstance(cv.value, bool) else "left"))
            if hor == "general": hor = "right" if isinstance(cv.value, (int, float)) else "left"
            envolver = bool(cf.alignment.wrap_text) or (isinstance(cv.value, str) and len(cv.value) > 60 and (r2 > r or c2 > c))
            pad = 3
            kw = dict(fontsize=tam * (0.92 if FUENTE in ("Calibri", "Carlito") else 0.84), fontweight=peso, fontstyle=estilo, color=col_txt, family=FUENTE, va="center")
            if signo:                                        # formato contable
                ax.text(x0 + pad, y0 + h / 2, signo, ha="left", **kw)
                ax.text(x0 + w - pad, y0 + h / 2, texto, ha="right", **kw)
                continue
            if envolver:
                import textwrap
                ancho_chars = max(10, int(w / (tam * 0.55)))
                lineas = []
                for parrafo in str(texto).split("\n"):
                    lineas += textwrap.wrap(parrafo, ancho_chars) or [""]
                caben = max(1, int(h / (tam * 1.45)))
                lineas = lineas[:caben]
                ax.text(x0 + pad if hor == "left" else (x0 + w / 2 if hor == "center" else x0 + w - pad),
                        y0 + pad + tam * 0.8, "\n".join(lineas), ha=hor, va="top", linespacing=1.15,
                        **{k: val for k, val in kw.items() if k != "va"})
                continue
            xs = x0 + pad if hor == "left" else (x0 + w / 2 if hor == "center" else x0 + w - pad)
            t = str(texto)
            maxc = int((w - 2 * pad) / (tam * 0.55))
            if len(t) > maxc and not isinstance(cv.value, (int, float)):
                # la celda de al lado esta ocupada: Excel recorta el texto
                vecina = wsv.cell(row=r, column=min(c2 + 1, c1)).value if c2 < c1 else None
                if vecina not in (None, ""):
                    t = t[:max(1, maxc)]
            ax.text(xs, y0 + h / 2, t, ha=hor, **kw)

    os.makedirs(os.path.dirname(salida), exist_ok=True)
    fig.savefig(salida, dpi=200, facecolor="white")
    plt.close(fig)
    return salida


if __name__ == "__main__":
    import sys
    rango(sys.argv[1], sys.argv[2], sys.argv[3])
